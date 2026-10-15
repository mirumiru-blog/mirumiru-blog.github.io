"""把 Excel 歌單轉成網頁使用的 site/data/songs.json。

用法：python scripts/build_songs.py [Excel 檔路徑]

資料來源：
- 「資料庫」工作表中的「所有歌曲」表格：歌曲名稱 / 歌曲語言 / 歌曲類別
- 各分類分頁（JPOP、アニソン、ボカロ、中文…）：歌名 / 可用連結
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "杏仁ミル_歌回歌單.xlsx"
OUTPUT = ROOT / "site" / "data" / "songs.json"

DB_SHEET = "資料庫"
DB_TABLE = "所有歌曲"
UNKNOWN = "未分類"

_INVISIBLE = re.compile("[​‌‍﻿]")


def clean(value):
    """去掉不可見字元與多餘空白，空值回傳空字串。"""
    if value is None:
        return ""
    text = _INVISIBLE.sub("", str(value)).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def match_key(name):
    """比對連結用的鍵：忽略空白與大小寫，讓「王心凌-  月光」對得上「王心凌- 月光」。"""
    return re.sub(r"\s+", "", clean(name)).lower()


def read_db_rows(wb):
    ws = wb[DB_SHEET]
    table = ws.tables.get(DB_TABLE)
    ref = table.ref if table else ws.dimensions
    rows = ws[ref]
    return [[cell.value for cell in row[:3]] for row in rows[1:]]


def read_links(wb):
    """從所有「歌名 / 可用連結」格式的分頁收集連結。"""
    links = {}
    for ws in wb.worksheets:
        header = [clean(c.value) for c in ws[1][:2]]
        if header != ["歌名", "可用連結"]:
            continue
        for name, url in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            url = clean(url)
            if clean(name) and url.startswith("http"):
                links.setdefault(match_key(name), url)
    return links


def build_songs(db_rows, links):
    songs = []
    seen = set()
    for row in db_rows:
        name, language, category = (clean(v) for v in (list(row) + [None] * 3)[:3])
        if not name:
            continue
        key = (match_key(name), language, category)
        if key in seen:
            continue
        seen.add(key)
        songs.append(
            {
                "name": name,
                "language": language or UNKNOWN,
                "category": category or UNKNOWN,
                "link": links.get(match_key(name)),
            }
        )
    return songs


def main(argv):
    xlsx = Path(argv[1]) if len(argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    songs = build_songs(read_db_rows(wb), read_links(wb))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        json.dumps(songs, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    with_link = sum(1 for s in songs if s["link"])
    print(f"已輸出 {len(songs)} 首歌（{with_link} 首有連結）到 {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main(sys.argv)
